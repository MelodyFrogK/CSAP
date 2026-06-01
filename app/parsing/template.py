"""KISA CSAP 명세서 엑셀을 읽어 점검항목과 시트별 열 배치를 파악한다.

실제 KISA 양식(SaaS 표준/간편 등)의 여러 레이아웃을 지원한다.
- 표준(17열): 분야·통제항목·(세부통제)·세부통제내용·관련법규·점검항목·해설(2024.6/2024.7)·
  운영여부·운영현황·관련문서·운영증적·담당자·점검결과·점검결과근거·심사위원·메모
- 간편/갱신: 평가기관 평가현황·최초 심사현황·운영현황 등이 분리됨(운영여부 열이 없는 시트도 있음)

헤더 문구로 열을 매핑한다. 점검항목 해설이 여러 버전이면 가장 최신(오른쪽) 열을 쓴다.
분야/통제항목/세부통제내용/관련법규는 블록 첫 행에만 채워지므로 아래 행으로 forward-fill 한다.
'점검항목' 열이 채워진 행만 작성 단위로 본다. 평가자 전용 열(점검결과·심사위원 등)은 무시한다.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from ..models import ControlItem, SheetLayout

# 안내 시트 등 명세 대상이 아닌 시트
SKIP_SHEETS = {"작성방법", "표지", "목차", "개요"}

_MAX_HEADER_SCAN = 12
# 데이터 시작 후 연속 빈 행이 이만큼 나오면 표 끝으로 간주(부풀려진 dims 방어)
_MAX_BLANK_STREAK = 80


def _norm(v: object) -> str:
    return str(v).replace(" ", "").replace("\n", "").lower() if v is not None else ""


def _classify(header: str) -> str | None:
    """헤더 셀 문구 → 필드. 더 긴 접두어를 먼저 본다('점검항목해설' > '점검항목')."""
    h = header
    if not h:
        return None
    # 결과/특수 열 (구체적인 것 우선)
    if h.startswith("운영여부"):
        return "operation"
    if h.startswith("운영현황"):
        return "status"
    if h.startswith("관련문서"):
        return "related_docs"
    if h.startswith("운영증적"):
        return "evidence"
    if h.startswith("증적확인담당자") or "담당자" in h:
        return "person"
    if h.startswith("보완"):
        return "improvement"
    # 입력 열
    if h.startswith("점검항목해설") or h == "해설":
        return "explanation"
    if h.startswith("점검항목"):
        return "check"
    if h.startswith("관련법규"):
        return "related_law"
    if h.startswith("세부통제내용") or h.startswith("세부통제") and "내용" in h:
        return "detail"
    if h.startswith("통제항목"):
        return "control"
    if h == "분야" or h.startswith("분야"):
        return "domain"
    return None


def detect_layout(ws: Worksheet) -> SheetLayout:
    """헤더 행을 찾아 열 배치를 추정한다."""
    best_row, best_score, best_map = 1, -1, {}
    for r in range(1, min(_MAX_HEADER_SCAN, ws.max_row) + 1):
        mapping: dict[str, int] = {}
        for c in range(1, ws.max_column + 1):
            field = _classify(_norm(ws.cell(row=r, column=c).value))
            if not field:
                continue
            # 점검항목 해설이 여러 버전(예: 2024.6/2024.7)이면 가장 최신(오른쪽) 열을 사용
            if field == "explanation":
                mapping["explanation"] = c
            elif field not in mapping:
                mapping[field] = c
        # 점검항목+운영현황이 동시에 있으면 강하게 가중
        score = len(mapping) + (5 if "check" in mapping and "status" in mapping else 0)
        if score > best_score:
            best_row, best_score, best_map = r, score, mapping

    layout = SheetLayout(
        sheet=ws.title,
        header_row=best_row,
        data_start_row=best_row + 1,
        col_domain=best_map.get("domain"),
        col_control=best_map.get("control"),
        col_detail=best_map.get("detail"),
        col_related_law=best_map.get("related_law"),
        col_check=best_map.get("check"),
        col_explanation=best_map.get("explanation"),
        col_operation=best_map.get("operation"),
        col_status=best_map.get("status"),
        col_related_docs=best_map.get("related_docs"),
        col_evidence=best_map.get("evidence"),
        col_person=best_map.get("person"),
        col_improvement=best_map.get("improvement"),
    )

    # 세부 통제항목 코드 열(헤더 없이 통제항목 헤더와 병합된 경우): 통제항목 바로 옆 빈 헤더 열
    mapped_cols = {c for c in best_map.values()}
    if layout.col_control:
        cand = layout.col_control + 1
        if cand <= ws.max_column and cand not in mapped_cols and cand != layout.col_detail:
            layout.col_sub_control = cand
    return layout


def read_sheet(ws: Worksheet) -> tuple[SheetLayout, list[ControlItem]]:
    layout = detect_layout(ws)
    items: list[ControlItem] = []
    if not layout.is_spec_sheet():
        return layout, items

    # forward-fill 상태
    last = {"domain": "", "control": "", "sub_control": "", "detail": "", "related_law": ""}

    # 실제 KISA 파일은 시트 dims가 100만 행까지 부풀려진 경우가 많다.
    # 데이터 시작 후 연속 빈 행이 일정 수 이상이면 표가 끝난 것으로 보고 중단한다.
    blank_streak = 0
    key_cols = [
        c for c in (
            layout.col_domain, layout.col_control, layout.col_sub_control,
            layout.col_detail, layout.col_check, layout.col_explanation,
        ) if c
    ]

    for r in range(layout.data_start_row, ws.max_row + 1):
        def cell(col: int | None) -> str:
            if not col:
                return ""
            v = ws.cell(row=r, column=col).value
            return str(v).strip() if v is not None else ""

        if not any(cell(c) for c in key_cols):
            blank_streak += 1
            if blank_streak >= _MAX_BLANK_STREAK:
                break
            continue
        blank_streak = 0

        # 컨텍스트 열 forward-fill
        for key, col in (
            ("domain", layout.col_domain),
            ("control", layout.col_control),
            ("sub_control", layout.col_sub_control),
            ("detail", layout.col_detail),
            ("related_law", layout.col_related_law),
        ):
            val = cell(col)
            if val:
                last[key] = val

        check = cell(layout.col_check)
        if not check:
            continue  # 점검항목이 없는 행은 작성 단위가 아님

        items.append(
            ControlItem(
                sheet=ws.title,
                row=r,
                domain=last["domain"],
                control=last["control"],
                sub_control=last["sub_control"],
                detail=last["detail"],
                related_law=last["related_law"],
                check_item=check,
                explanation=cell(layout.col_explanation),
            )
        )
    return layout, items


def read_template(path: str | Path) -> tuple[list[ControlItem], dict[str, SheetLayout]]:
    """모든 명세 시트에서 점검항목과 시트별 레이아웃을 읽는다."""
    wb = load_workbook(str(path), data_only=True)
    items: list[ControlItem] = []
    layouts: dict[str, SheetLayout] = {}
    for ws in wb.worksheets:
        if ws.title.strip() in SKIP_SHEETS:
            continue
        layout, sheet_items = read_sheet(ws)
        if not layout.is_spec_sheet():
            continue
        layouts[ws.title] = layout
        items.extend(sheet_items)
    return items, layouts
