"""분석 결과를 원본 KISA 명세서에 써넣어 새 파일로 저장한다.

원본 서식을 보존하기 위해 템플릿을 그대로 로드한다.
- 기존 폼: 운영여부/운영현황/관련문서/운영증적 열에 그대로 채운다.
- 추가 폼: 템플릿에 없는 보조 열(자동 판정·판단 근거·보완사항·준비 권장 증적·검토 필요)을
  맨 오른쪽(평가자 열 뒤)에 새로 추가한다.
증적확인 담당자 열은 사람이 채워야 하므로 건드리지 않는다.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ..models import ItemResult, SheetLayout

# 오른쪽에 추가하는 자동작성 열 정의: (layout 속성, ItemResult 속성, 헤더, 너비)
EXTRA_COLUMNS = [
    ("col_assessment", "assessment", "자동 판정(상세)\n[참고용]", 14),
    ("col_basis", "basis", "판단 근거\n(해설 요건 대조)", 38),
    ("col_improvement", "improvement", "보완사항(자동작성)", 46),
    ("col_recommended", "recommended", "준비 권장 증적·문서\n(색인 예: 1.1.1.1-1)", 38),
    ("col_review", "review_needed", "검토 필요", 10),
]

_HEADER_FILL = PatternFill("solid", fgColor="FFF2CC")  # 연노랑(추가 열 식별용)
_HEADER_FONT = Font(bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")
_HEADER_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")


def write_results(
    template_path: str | Path,
    layouts: dict[str, SheetLayout],
    results: list[ItemResult],
    output_path: str | Path,
) -> Path:
    wb = load_workbook(str(template_path))

    # 시트별로 추가 열을 (없으면) 끝에 만든다.
    for sheet, layout in layouts.items():
        ws = wb[sheet]
        next_col = ws.max_column + 1
        for attr, _res_attr, header, width in EXTRA_COLUMNS:
            if getattr(layout, attr) is not None:
                continue  # 템플릿에 이미 있으면 재사용
            col = next_col
            next_col += 1
            setattr(layout, attr, col)
            hcell = ws.cell(row=layout.header_row, column=col, value=header)
            hcell.font = _HEADER_FONT
            hcell.fill = _HEADER_FILL
            hcell.alignment = _HEADER_ALIGN
            ws.column_dimensions[hcell.column_letter].width = width

    for res in results:
        layout = layouts.get(res.sheet)
        if not layout:
            continue
        ws = wb[res.sheet]
        # 기존 폼
        _set(ws, res.row, layout.col_operation, res.operation)
        _set(ws, res.row, layout.col_status, res.status)
        _set(ws, res.row, layout.col_related_docs, res.related_docs)
        _set(ws, res.row, layout.col_evidence, res.evidence)
        # 추가 폼
        for attr, res_attr, _header, _width in EXTRA_COLUMNS:
            _set(ws, res.row, getattr(layout, attr), getattr(res, res_attr))

    # 일부 실제 KISA 파일은 빈 셀이 100만 행까지 서식 지정돼 있어 저장이 매우 느리고
    # 결과 파일도 비대해진다. 데이터 마지막 행 이후의 유령 셀을 정리한다.
    last_rows: dict[str, int] = {}
    for res in results:
        last_rows[res.sheet] = max(last_rows.get(res.sheet, 0), res.row)
    for sheet, layout in layouts.items():
        keep = max(last_rows.get(sheet, layout.data_start_row), layout.header_row) + 1
        _prune_phantom_rows(wb[sheet], keep)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    return output_path


def _set(ws, row: int, col: int | None, value: str) -> None:
    if not col or value == "":
        return
    cell = ws.cell(row=row, column=col, value=value)
    cell.alignment = _WRAP


def _prune_phantom_rows(ws, keep_through: int) -> None:
    """keep_through 행 이후에 남은 (대개 빈) 셀을 제거해 저장 속도/파일 크기를 개선한다.

    부풀려진 dims(예: 1,048,386행)를 가진 파일에서 wb.save 가 수십 초~분 걸리는 문제를
    완화한다. openpyxl 내부 셀 저장소를 직접 정리하고 dimension 캐시를 재계산한다.
    """
    cells = getattr(ws, "_cells", None)
    if not cells:
        return
    if ws.max_row <= keep_through:
        return
    stale = [coord for coord in cells if coord[0] > keep_through]
    for coord in stale:
        del cells[coord]
    # row_dimensions 도 정리
    for r in [r for r in ws.row_dimensions if r > keep_through]:
        del ws.row_dimensions[r]
    ws._current_row = min(ws._current_row, keep_through) if hasattr(ws, "_current_row") else keep_through
