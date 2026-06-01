"""핵심 파이프라인 스모크 테스트 (echo 공급자, 키 불필요)."""
from __future__ import annotations

import os
from pathlib import Path

os.environ.setdefault("CSAP_LLM_PROVIDER", "echo")

from openpyxl import load_workbook  # noqa: E402

from app.core import pipeline  # noqa: E402
from app.core.retrieval import split_into_chunks, top_chunks  # noqa: E402
from app.llm.base import LLMProvider  # noqa: E402
from app.parsing import documents, template  # noqa: E402
from app.parsing.template import _classify  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
SAMPLE_TEMPLATE = ROOT / "data" / "sample_csap_template.xlsx"


def _ensure_template() -> Path:
    if not SAMPLE_TEMPLATE.exists():
        from scripts.make_sample_template import main as make

        make()
    return SAMPLE_TEMPLATE


def test_extract_json_variants():
    p = LLMProvider
    assert p._extract_json('{"a": 1}') == {"a": 1}
    assert p._extract_json('```json\n{"a": 2}\n```') == {"a": 2}
    assert p._extract_json('설명입니다.\n{"a": 3}\n끝.') == {"a": 3}


def test_header_classify():
    # 접두어 충돌: '점검항목해설' 이 '점검항목'보다 먼저 분류돼야 함
    assert _classify("점검항목해설") == "explanation"
    assert _classify("점검항목") == "check"
    assert _classify("통제항목") == "control"
    assert _classify("운영여부") == "operation"
    assert _classify("운영현황") == "status"
    assert _classify("관련문서(정책,지침등세부조항번호까지)") == "related_docs"
    assert _classify("운영증적(자산,파일등)") == "evidence"
    assert _classify("증적확인담당자(소속,이름,연락처)") == "person"
    assert _classify("관련법규") == "related_law"
    # 평가자 전용 열은 작성 대상이 아님(분류되지 않아야 함)
    assert _classify("점검결과") is None
    assert _classify("점검결과근거") is None
    assert _classify("심사위원") is None


def test_chunking_and_retrieval():
    text = "접근 통제 정책. 사용자 계정은 최소권한으로 관리한다. " * 50
    chunks = split_into_chunks(text, "doc.txt", size=200, overlap=40)
    assert chunks
    assert top_chunks("계정 관리 최소권한 접근 통제", chunks, k=3)


def test_template_detection_and_forward_fill():
    items, layouts = template.read_template(_ensure_template())
    assert len(layouts) == 2  # 관리적 / 기술적
    assert len(items) >= 5
    # 점검항목 단위로 분해되고 컨텍스트가 forward-fill 되는지
    multi = [i for i in items if i.check_item.startswith("2)")]
    assert multi and all(i.domain and i.control for i in multi)
    layout = next(iter(layouts.values()))
    assert layout.col_check and layout.col_status and layout.col_operation
    # 실제 양식 기반: 관련 법규 열과 (해설 2개 버전 중) 최신 해설 열을 인식
    assert layout.col_related_law and layout.col_explanation
    # 점검항목 해설(2024.7) = 8열(H)을 최신본으로 선택
    assert layout.col_explanation == 8


def test_full_pipeline_echo(tmp_path):
    tpl = _ensure_template()
    doc = tmp_path / "policy.txt"
    doc.write_text(
        "당사는 정보보호 정책을 수립하고 정보보호 최고책임자(CISO)의 승인을 받았다. "
        "사용자 계정 생성·변경·삭제 절차를 운영하고 관리자 접근에 다중인증(MFA)을 적용한다.",
        encoding="utf-8",
    )
    out = tmp_path / "result.xlsx"
    res = pipeline.run(str(tpl), [str(doc)], str(out))
    assert out.exists()
    assert res.processed == res.item_count >= 5
    assert not res.errors

    # 결과 열(운영현황/보완사항)이 실제로 채워졌는지 확인
    _, layouts = template.read_template(out)
    wb = load_workbook(str(out))
    layout = layouts["1. 관리적 보호조치"]
    assert layout.col_improvement is not None  # 보완사항 열 추가됨
    ws = wb["1. 관리적 보호조치"]
    filled = sum(
        1
        for r in range(layout.data_start_row, ws.max_row + 1)
        if ws.cell(row=r, column=layout.col_status).value
        or ws.cell(row=r, column=layout.col_improvement).value
    )
    assert filled > 0


def test_txt_extraction(tmp_path):
    f = tmp_path / "policy.txt"
    f.write_text("정보보호 정책을 수립하고 경영진 승인을 받았다.", encoding="utf-8")
    assert "정보보호" in documents.extract_text(f)


def test_system_prompt_has_antihallucination():
    from app.core import prompts

    # 주제 환각 방지 규칙과 판정 기준이 프롬프트에 포함돼야 함
    assert "환각" in prompts.SYSTEM_PROMPT
    assert "확인불가" in prompts.SYSTEM_PROMPT
    assert "공급망" in prompts.SYSTEM_PROMPT  # 반례 예시


def test_echo_distinguishes_relevance():
    """근거 있는 항목은 운영, 무관한 항목은 확인불가로 구분되어야 한다."""
    from app.core import prompts
    from app.llm.echo_provider import EchoProvider
    from app.models import ControlItem

    p = EchoProvider()
    excerpt_doc = "당사는 정보보호 정책을 수립하고 CISO 승인 하에 지침을 문서화하여 운영한다."

    related = ControlItem(
        sheet="s", row=5, check_item="정보보호 정책을 수립하고 지침을 문서화하고 있는가?"
    )
    unrelated = ControlItem(
        sheet="s", row=6, check_item="가상자원의 생성·변경·회수 절차를 운영하는가?"
    )

    from app.core.retrieval import Chunk

    chunk = [Chunk(doc="d", index=0, text=excerpt_doc)]
    r1 = p.complete_json(prompts.SYSTEM_PROMPT, prompts.build_user_prompt(related, chunk))
    r2 = p.complete_json(prompts.SYSTEM_PROMPT, prompts.build_user_prompt(unrelated, chunk))
    assert r1["assessment"] == "운영" and r1["operation"] == "Y"
    assert r2["assessment"] == "확인불가" and r2["review_needed"] == "Y"
